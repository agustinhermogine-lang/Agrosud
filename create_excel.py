from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from datetime import datetime
from pathlib import Path
out=Path('NABSA_MAIZ_SAILED_FECHAS_CORRECTAS_2026-09-10.xlsx')
R='09/09/2026'
# puerto, muelle/terminal, buque, fecha zarpe, toneladas, destino
rows=[
('SAN LORENZO','COFCO INTL. PGSM SOUTH BERTH (EX NIDERA SAN MARTIN)','MALLIKA NAREE','01/09/2026','16.999,86','CONGO'),('SAN LORENZO','COFCO INTL. PGSM SOUTH BERTH (EX NIDERA SAN MARTIN)','MALLIKA NAREE','01/09/2026','10.500,00','ANGOLA'),('SAN LORENZO','PAMPA SAN MARTIN','BERGE ASAHIDAKE','01/09/2026','33.000,00','TUNISIA'),('SAN LORENZO','TERMINAL 6 SOUTH BERTH SAN MARTIN','DSM HARBOUR','01/09/2026','18.137,62','YEMEN'),
('SAN LORENZO','A.C.A. TIMBUES','FORTUNE KNIGHT','02/09/2026','49.499,85','INDONESIA'),('SAN LORENZO','QUEBRACHO SAN MARTIN','BERGE JUNGFRAU','03/09/2026','14.877,87','CHILE'),('SAN LORENZO','A.C.A. TIMBUES','AGIOS IOANNIS','04/09/2026','29.934,29','MALAYSIA'),('SAN LORENZO','COFCO INTL. SOUTH BERTH (EX COFCO SOUTH)','IRVINE BAY','04/09/2026','22.618,16','PERU'),('SAN LORENZO','ACA SAN LORENZO','LEFKADA I','05/09/2026','49.522,00','VIETNAM'),('SAN LORENZO','RENOVA NORTH','ISKENDERUN-M','05/09/2026','30.000,00','EGYPT'),('SAN LORENZO','QUEBRACHO SAN MARTIN','DRINA S','05/09/2026','19.784,73','MOROCCO'),('SAN LORENZO','TERMINAL 6 SOUTH BERTH SAN MARTIN','AGIA PARASKEVI','06/09/2026','15.155,50','MOROCCO'),('SAN LORENZO','SAN BENITO SAN LORENZO','ARIETTA','06/09/2026','26.846,26','MALAYSIA'),('SAN LORENZO','COFCO INTL. PGSM SOUTH BERTH (EX NIDERA SAN MARTIN)','STAR WESTPORT','07/09/2026','36.000,34','PHILIPPINES'),('SAN LORENZO','TRANSITO (ADM AGRO)','RANA','08/09/2026','35.904,00','CHILE'),('SAN LORENZO','QUEBRACHO SAN MARTIN','YAMI BENEFIT','09/09/2026','33.286,19','CHILE'),('SAN LORENZO','A.G.D. TIMBUES','INGENUITY','09/09/2026','46.685,82','SAUDI ARABIA'),
('ROSARIO','CARGILL PUNTA ALVEAR','CENTURION SIGNIFER','01/09/2026','22.928,96','PERU'),('ROSARIO','CARGILL PUNTA ALVEAR','TOKYO SPIRIT','03/09/2026','3.000,00','PERU'),('ROSARIO','CARGILL PUNTA ALVEAR','TOKYO SPIRIT','03/09/2026','17.348,97','PERU'),('ROSARIO','LDC TERMINAL GRAL. LAGOS','LIMNIONAS','04/09/2026','9.000,00','VIETNAM'),('ROSARIO','LDC TERMINAL GRAL. LAGOS','LIMNIONAS','04/09/2026','47.562,21','VIETNAM'),('ROSARIO','CARGILL PUNTA ALVEAR','SOFIA','07/09/2026','35.521,99','PERU'),('ROSARIO','ADM AGRO ARROYO SECO (EX TOEPFER)','CL KIBOU','07/09/2026','35.200,00','PERU'),('ROSARIO','ADM AGRO ARROYO SECO (EX TOEPFER)','MEDITERRANEAN SPIRIT','09/09/2026','35.200,00','PERU'),
('LIMA','DELTA DOCK LIMA','DIAMOND SKY','04/09/2026','38.000,00','PERU'),('NECOCHEA / QUEQUEN','BERTH 1 OPEN QUAY (QUEQUEN) (SITIO 0 DE QUEQUEN TERMINAL)','GALAPAGOS','04/09/2026','24.317,00','VIETNAM'),
('BAHIA BLANCA','CARGILL TERMINAL PUERTO INGENIERO WHITE','ARTUNIS','01/09/2026','18.680,00','SAUDI ARABIA'),('BAHIA BLANCA','LDC GRAIN TERMINAL PUERTO GALVAN','CIC EPOS','03/09/2026','24.133,00','MALAYSIA'),('BAHIA BLANCA','CARGILL TERMINAL PUERTO INGENIERO WHITE','IONIC KLEOS','04/09/2026','24.335,00','VIETNAM'),('BAHIA BLANCA','LDC GRAIN TERMINAL PUERTO GALVAN','TAHO EUDAIMONIA','06/09/2026','10.319,00','MALAYSIA'),('BAHIA BLANCA','TERMINAL BAHIA BLANCA BERTHS 9 PUERTO INGENIERO WHITE','PHAEDRA','08/09/2026','23.812,84','MALAYSIA'),('BAHIA BLANCA','BERTHS 2 / 3 BUNGE (EX VITERRA EX MORENO) PUERTO GALVÁN','DIANE','08/09/2026','16.000,00','SAUDI ARABIA'),('BAHIA BLANCA','TERMINAL BAHIA BLANCA BERTHS 9 PUERTO INGENIERO WHITE','DIANE','08/09/2026','6.705,00','SAUDI ARABIA')]
wb=Workbook(); ws=wb.active; ws.title='Detalle'
headers=['Puerto','Muelle/Terminal','Buque','Fecha de zarpe','Toneladas','Cargo','Origin','Destino','Fecha del reporte']
ws.append(headers)
for r in rows:
 d=datetime.strptime(r[3],'%d/%m/%Y'); t=float(r[4].replace('.','').replace(',','.'))
 ws.append([r[0],r[1],r[2],d,t,'CORN','ARGENTINA',r[5],datetime.strptime(R,'%d/%m/%Y')])
for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='1F4E78')
for row in ws.iter_rows(min_row=2): row[3].number_format=row[8].number_format='dd/mm/yyyy'; row[4].number_format='#,##0.00'
for sh, key in [('Total por puerto',0),('Total por muelle',1)]:
 s=wb.create_sheet(sh); s.append([headers[key],'Registros','Toneladas'])
 agg={}
 for r in rows: agg.setdefault(r[key],[0,0]); agg[r[key]][0]+=1; agg[r[key]][1]+=float(r[4].replace('.','').replace(',','.'))
 for k,(n,total) in sorted(agg.items()): s.append([k,n,total]); s.cell(s.max_row,3).number_format='#,##0.00'
 for c in s[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E78')
s=wb.create_sheet('Fuentes'); s.append(['Elemento','Detalle']); info=[('Consulta','10/09/2026'),('Reporte público','09/09/2026'),('URL','https://www.nabsa.com.ar/assets/vessels_sailed_update.pdf'),('TXT utilizado',r'C:\Users\Usuario\AppData\Local\Temp\nabsa-pdf-e7d782b4273f43ed944456a176a7f21d\vessels_sailed_update.txt'),('Filtro','Cargo=CORN; Origin=ARGENTINA; fechas de zarpe 01/09/2026 a 10/09/2026 inclusive'),('Aclaración','No había reporte del 10/09/2026; se utilizó el reporte más reciente del 09/09/2026.')]
for x in info:s.append(x)
for c in s[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E78')
for sh in wb.worksheets:
 sh.freeze_panes='A2'; sh.auto_filter.ref=sh.dimensions
 for col in range(1,sh.max_column+1): sh.column_dimensions[get_column_letter(col)].width=min(60,max(14,max(len(str(sh.cell(r,col).value or '')) for r in range(1,sh.max_row+1))+2))
wb.save(out)
# validation with openpyxl
v=load_workbook(out,data_only=True); d=v['Detalle']; dates=[d.cell(r,4).value for r in range(2,d.max_row+1)]; total=sum(d.cell(r,5).value for r in range(2,d.max_row+1))
assert v.sheetnames==['Detalle','Total por puerto','Total por muelle','Fuentes'] and len(dates)==34 and min(dates).strftime('%d/%m/%Y')=='01/09/2026' and max(dates).strftime('%d/%m/%Y')=='09/09/2026'
print(f'CREADO {out.resolve()} | hojas={v.sheetnames} | registros={len(dates)} | minimo={min(dates):%d/%m/%Y} | maximo={max(dates):%d/%m/%Y} | total={total:,.2f}')
